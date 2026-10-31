"""Exercise real calendar boundaries and weekly scope through the intake APIs."""

from dataclasses import replace
from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.demo.generate import workbook
from app.models import AttendanceCase, Email
from app.services.imports import process_email, synchronize
from app.services.review_calendar import business_today, default_cycle, payroll_period, period_bounds, period_schedule
from tests.helpers import Provider, ingest, row


@pytest.mark.parametrize(
    ("day", "period"),
    [
        (date(2026, 9, 26), "2026-10"),
        (date(2026, 10, 25), "2026-10"),
        (date(2026, 10, 26), "2026-11"),
        (date(2026, 10, 30), "2026-11"),
        (date(2026, 12, 25), "2026-12"),
        (date(2026, 12, 26), "2027-01"),
        (date(2028, 2, 29), "2028-03"),
    ],
)
def test_payroll_ownership_rolls_over_on_the_26th(day, period):
    assert payroll_period(day) == period
    starts, ends = period_bounds(period)
    assert starts <= day <= ends
    assert starts.day == 26 and ends.day == 25


def test_payroll_migration_updates_existing_periods_and_can_restore_calendar_bounds(tmp_path, monkeypatch):
    from pathlib import Path

    from alembic.config import Config
    from sqlalchemy import create_engine, text

    from alembic import command

    database_url = "sqlite:///" + str(tmp_path / "migration.sqlite")
    monkeypatch.setenv("DATABASE_URL", database_url)
    root = Path(__file__).resolve().parents[1]
    config = Config(str(root / "alembic.ini"))
    config.set_main_option("script_location", str(root / "alembic"))
    command.upgrade(config, "a6c832f03190")
    engine = create_engine(database_url)
    with engine.begin() as connection:
        connection.execute(
            text(
                "INSERT INTO review_periods (id,name,starts,ends) VALUES ('2028-03','Historic period','2028-03-01','2028-03-31')"
            )
        )
    command.upgrade(config, "head")
    with engine.connect() as connection:
        row = connection.execute(text("SELECT id,name,starts,ends FROM review_periods")).one()
        assert tuple(row) == ("2028-03", "Historic period", "2028-02-26", "2028-03-25")
    command.downgrade(config, "a6c832f03190")
    with engine.connect() as connection:
        row = connection.execute(text("SELECT id,name,starts,ends FROM review_periods")).one()
        assert tuple(row) == ("2028-03", "Historic period", "2028-03-01", "2028-03-31")
    engine.dispose()


@pytest.mark.parametrize(
    ("today", "expected"),
    [
        (date(2026, 10, 25), "2026-10"),
        (date(2026, 10, 26), "2026-11"),
        (date(2026, 10, 30), "2026-11"),
        (date(2026, 12, 26), "2027-01"),
    ],
)
def test_bootstrap_selects_current_payroll_period_at_cutoff(context, monkeypatch, today, expected):
    from app.services import catalog

    _, client = context
    monkeypatch.setattr(catalog, "business_today", lambda: today)
    response = client.get("/api/bootstrap")
    assert response.status_code == 200
    data = response.json()
    assert data["current_period"] == expected
    selected = next(period for period in data["periods"] if period["id"] == expected)
    assert selected["starts"] == period_bounds(expected)[0].isoformat()
    assert selected["ends"] == period_bounds(expected)[1].isoformat()
    assert selected["name"].endswith(f"{int(expected[5:]):02}/{expected[:4]}")


@pytest.mark.parametrize(
    ("period", "saturdays", "mondays", "last_due"),
    [
        ("2026-10", ["2026-09-26", "2026-10-03", "2026-10-10", "2026-10-17", "2026-10-24"], 4, "2026-10-26"),
        ("2026-11", ["2026-10-31", "2026-11-07", "2026-11-14", "2026-11-21"], 5, "2026-11-23"),
        ("2021-02", ["2021-01-30", "2021-02-06", "2021-02-13", "2021-02-20"], 4, "2021-02-22"),
        ("2020-03", ["2020-02-29", "2020-03-07", "2020-03-14", "2020-03-21"], 4, "2020-03-23"),
        ("2027-01", ["2026-12-26", "2027-01-02", "2027-01-09", "2027-01-16", "2027-01-23"], 5, "2027-01-25"),
    ],
)
def test_real_month_lengths_and_cross_month_deadlines(period, saturdays, mondays, last_due):
    schedule = period_schedule(period, date(2026, 10, 30))
    assert [cycle["send_on"].isoformat() for cycle in schedule["cycles"]] == saturdays
    assert schedule["saturday_count"] == len(saturdays)
    assert schedule["monday_count"] == mondays
    assert schedule["cycles"][-1]["due_on"].isoformat() == last_due
    assert all(cycle["send_on"].weekday() == 5 and cycle["due_on"].weekday() == 0 for cycle in schedule["cycles"])


def test_business_day_uses_vietnam_timezone_at_midnight():
    today = business_today(datetime(2026, 11, 1, 17, 0, tzinfo=timezone.utc))
    assert today == date(2026, 11, 2)
    assert period_schedule("2026-11", today)["cycles"][0]["state"] == "DUE_TODAY"
    with pytest.raises(ValueError):
        business_today(datetime(2026, 11, 1))


def test_next_payroll_period_links_to_the_previous_periods_closing_monday():
    november = period_schedule("2026-11", date(2026, 11, 1))
    assert november["saturday_count"] == 4 and november["monday_count"] == 5
    assert november["carryover"] == {
        "period_id": "2026-10",
        "send_on": date(2026, 10, 24),
        "due_on": date(2026, 10, 26),
    }
    assert period_schedule("2026-10", date(2026, 10, 30))["carryover"] is None


@pytest.mark.parametrize(
    ("today", "expected"),
    [
        (date(2026, 10, 30), date(2026, 10, 24)),
        (date(2026, 10, 4), date(2026, 10, 3)),
        (date(2026, 10, 5), date(2026, 10, 3)),
        (date(2026, 10, 6), date(2026, 10, 10)),
    ],
)
def test_default_cycle_retains_open_windows_then_selects_the_next_saturday(today, expected):
    assert default_cycle("2026-10", today) == expected


def test_bootstrap_tracks_the_real_year_without_removing_historic_months(context, monkeypatch):
    from app.services import catalog

    _, client = context
    monkeypatch.setattr(catalog, "business_today", lambda: date(2027, 1, 1))
    data = client.get("/api/bootstrap").json()
    assert data["current_period"] == "2027-01"
    assert data["today"] == "2027-01-01"
    assert data["timezone"] == "Asia/Ho_Chi_Minh"
    assert {f"2027-{month:02}" for month in range(1, 13)} <= {period["id"] for period in data["periods"]}


class WeeklyProvider(Provider):
    def __init__(self, content, source, period, cycle):
        super().__init__(content, source)
        self.period, self.cycle = period, cycle

    def messages(self):
        return [
            replace(
                incoming,
                subject=f"[DCC] [{self.period}] [WEEK:{self.cycle}] Synthetic summary",
                received_at=datetime(2027, 1, 15, tzinfo=timezone.utc),
            )
            for incoming in super().messages()
        ]


def import_week(db, code, cycle, *, period="2026-10", day="2026-10-01"):
    provider = WeeklyProvider(workbook([row(code=code, day=day)]), code, period, cycle)
    synchronize(db, provider)
    email = db.scalar(select(Email).where(Email.source_id == code))
    process_email(db, email.id)
    db.commit()
    return email


@pytest.mark.parametrize(
    ("period", "cycle", "day", "accepted"),
    [
        ("2026-10", "2026-09-26", "2026-09-26", True),
        ("2026-10", "2026-10-24", "2026-10-25", True),
        ("2026-10", "2026-09-26", "2026-09-25", False),
        ("2026-10", "2026-10-24", "2026-10-26", False),
        ("2026-11", "2026-10-31", "2026-10-26", True),
        ("2027-01", "2026-12-26", "2026-12-26", True),
    ],
)
def test_intake_enforces_inclusive_payroll_boundaries(context, period, cycle, day, accepted):
    sessions, client = context
    with sessions() as db:
        email = import_week(db, "PAYROLL-BOUNDARY", cycle, period=period, day=day)
        case = db.scalar(select(AttendanceCase))
        assert (case is not None) == accepted
        assert email.status == ("IMPORTED" if accepted else "FAILED")
    listed = client.get("/api/cases", params={"period": period, "cycle": cycle}).json()
    assert listed["total"] == int(accepted)


def test_week_filters_agree_across_inbox_cases_dashboard_and_export(context):
    sessions, client = context
    with sessions() as db:
        legacy, _ = ingest(db, workbook([row(code="LEGACY")]))
        first = import_week(db, "WEEK-ONE", "2026-10-03")
        second = import_week(db, "WEEK-FIVE", "2026-10-24")
        assert legacy.cycle_start is None
        assert first.cycle_start == date(2026, 10, 3)
        assert second.cycle_start == date(2026, 10, 24)
    for cycle, code in [("2026-10-03", "WEEK-ONE"), ("2026-10-24", "WEEK-FIVE"), ("unassigned", "LEGACY")]:
        scope = {"period": "2026-10", "cycle": cycle}
        cases = client.get("/api/cases", params=scope)
        assert cases.status_code == 200
        assert [case["code"] for case in cases.json()["items"]] == [code]
        assert len(client.get("/api/emails", params=scope).json()) == 1
        dashboard = client.get("/api/dashboard", params=scope).json()
        assert dashboard["total"] == dashboard["emails"] == dashboard["departments_sent"] == 1
        export = client.get("/api/export", params=scope)
        book = load_workbook(BytesIO(export.content), read_only=True)
        assert book.active.max_row == 2
        assert book.active.cell(2, 1).value == code
        assert book.active.cell(2, 21).value == (None if cycle == "unassigned" else cycle)
        if cycle == "2026-10-24":
            assert book.active.cell(2, 22).value == "2026-10-26"
    assert client.get("/api/cases?period=2026-10").json()["total"] == 3
    assert client.get("/api/cases?period=2026-11&cycle=2026-10-31").json()["total"] == 0
    assert client.get("/api/cases?cycle=2026-10-99").status_code == 422
    assert client.get("/api/emails?cycle=invalid").status_code == 422


@pytest.mark.parametrize("cycle", ["2026-10-30", "2026-10-31", "2026-11-07", "2026-10-99", "invalid"])
def test_invalid_or_mismatched_week_requires_review_instead_of_month_fallback(context, cycle):
    sessions, _ = context
    with sessions() as db:
        email = import_week(db, "INVALID-CYCLE", cycle)
        assert email.status == "NEEDS_REVIEW"
        assert email.cycle_start is None
        assert db.scalar(select(AttendanceCase)) is None


def test_intake_can_create_a_new_calendar_month_and_keeps_late_replies_in_their_cycle(context):
    sessions, client = context
    with sessions() as db:
        email = import_week(db, "NEW-YEAR", "2028-01-22", period="2028-01", day="2028-01-20")
        assert email.period_id == "2028-01" and email.cycle_start == date(2028, 1, 22)
    assert client.get("/api/cases?period=2028-01&cycle=2028-01-22").json()["total"] == 1


def test_demo_jobs_validate_and_carry_the_selected_month_and_cycle(context):
    _, client = context
    assert client.post("/api/demo", json={"period": "2026-10", "cycle": "2026-10-30"}).status_code == 422
    response = client.post("/api/demo", json={"period": "2026-10", "cycle": "2026-10-24"})
    assert response.status_code == 200
    job = next(job for job in client.get("/api/jobs").json() if job["id"] == response.json()["job_id"])
    assert job["payload"] == {"period_id": "2026-10", "cycle_start": "2026-10-24"}


def test_calendar_demo_goes_through_intake_without_real_email_or_fixed_demo_month(context, monkeypatch):
    from app.demo import generate
    from app.services.email_provider import IncomingEmail

    sessions, _ = context
    captured = []

    class LocalSMTP:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def send_message(self, outgoing):
            captured.append(outgoing)

    monkeypatch.setattr(generate.smtplib, "SMTP", lambda *args, **kwargs: LocalSMTP())
    generate.run(period_id="2028-03", cycle_start="2028-02-26")
    assert len(captured) == 10
    assert "MP-2028-02-26-1-001" in str(captured[6]["Subject"])
    incoming = [
        IncomingEmail(
            source_id=str(index),
            message_id=str(mail["Message-ID"]),
            sender=str(mail["From"]),
            recipient=str(mail["To"]),
            subject=str(mail["Subject"]),
            body="",
            received_at=datetime(2028, 3, 15, tzinfo=timezone.utc),
            attachments=[(part.get_filename(), part.get_payload(decode=True)) for part in mail.iter_attachments()],
        )
        for index, mail in enumerate(captured)
    ]

    class DemoMailbox:
        def messages(self):
            return incoming

    with sessions() as db:
        assert synchronize(db, DemoMailbox())["new"] == 10
        for email in db.scalars(select(Email).order_by(Email.id)):
            process_email(db, email.id)
        db.commit()
        cases = list(db.scalars(select(AttendanceCase)))
        assert len(cases) == 60
        assert all(case.period_id == "2028-03" and case.cycle_start == date(2028, 2, 26) for case in cases)
        assert all(case.code.startswith("MP-2028-02-26-") for case in cases)
        assert synchronize(db, DemoMailbox())["existing"] == 10
