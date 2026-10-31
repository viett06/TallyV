"""Microsoft authentication and inbox ingestion without calling real services."""

import json
import logging
from datetime import timedelta
from email.message import EmailMessage
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

import httpx
import pytest
from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import select

from app.core.attendance import PUNCH_IN
from app.core.config import settings
from app.core.errors import AppError
from app.core.i18n import message
from app.core.middleware import OAuthAccessLogFilter
from app.core.time import utcnow
from app.demo.generate import workbook
from app.models import AttendanceCase, Email, OutlookAuthFlow, OutlookConnection, ProcessingJob
from app.services import mailbox, outlook_auth
from app.services.imports import process_email
from app.services.outlook_provider import OutlookProvider, graph_url
from tests.helpers import ingest, row

TENANT = "11111111-1111-4111-8111-111111111111"
CLIENT = "22222222-2222-4222-8222-222222222222"


@pytest.fixture
def microsoft(context, monkeypatch):
    monkeypatch.setattr(settings, "mail_provider", "outlook")
    monkeypatch.setattr(settings, "outlook_transport", "microsoft")
    monkeypatch.setattr(settings, "outlook_tenant_id", TENANT)
    monkeypatch.setattr(settings, "outlook_client_id", CLIENT)
    monkeypatch.setattr(settings, "outlook_client_secret", SecretStr("test-client-secret"))
    monkeypatch.setattr(settings, "outlook_test_mode", False)
    return context


class FakeIdentity:
    def __init__(self, cache):
        self.cache = cache

    def initiate_auth_code_flow(self, scopes, **kwargs):
        assert scopes == ["User.Read", "Mail.Read"]
        return {
            "state": "random-test-state",
            "auth_uri": "https://login.microsoftonline.com/authorize",
            "code_verifier": "private-pkce-verifier",
        }

    def acquire_token_by_auth_code_flow(self, flow, parameters):
        assert parameters["code"] == "test-auth-code"
        self.cache.deserialize(json.dumps({"test_refresh_token": "private-refresh-token"}))
        return {"access_token": "private-access-token", "id_token_claims": {"tid": TENANT}}

    def get_accounts(self):
        return [{"home_account_id": "test-account"}]


@pytest.fixture
def identity(monkeypatch):
    monkeypatch.setattr(outlook_auth, "microsoft_client", lambda cache: FakeIdentity(cache))
    monkeypatch.setattr(outlook_auth, "read_profile", lambda token: {"id": "mailbox-1", "mail": "hr@hospital.example"})


def test_authentication_state_is_browser_bound_one_time_and_tokens_remain_encrypted(microsoft, identity):
    sessions, client = microsoft
    with sessions() as db:
        url, cookie = outlook_auth.begin_connection(db, 1, "http://localhost:8000")
        assert url.startswith("https://login.microsoftonline.com/")
        flow = db.scalar(select(OutlookAuthFlow))
        assert "private-pkce-verifier" not in flow.payload
        assert "private-pkce-verifier" in outlook_auth.decrypt(flow.payload)
        with pytest.raises(AppError):
            outlook_auth.complete_connection(
                db, "wrong-browser-cookie", {"state": "random-test-state", "code": "test-auth-code"}
            )
        with pytest.raises(AppError):
            outlook_auth.complete_connection(db, cookie, {"state": "wrong-state", "code": "test-auth-code"})
        outlook_auth.complete_connection(db, cookie, {"state": "random-test-state", "code": "test-auth-code"})
        connection = db.get(OutlookConnection, 1)
        assert connection.mailbox == "hr@hospital.example" and connection.status == "CONNECTED"
        assert "private-refresh-token" not in connection.token_cache
        assert "private-refresh-token" in outlook_auth.decrypt(connection.token_cache)
        with pytest.raises(AppError):
            outlook_auth.complete_connection(db, cookie, {"state": "random-test-state", "code": "test-auth-code"})
    response = client.get("/api/outlook")
    assert response.status_code == 200
    assert all(
        value not in response.text
        for value in ["private-refresh-token", "private-access-token", "test-client-secret", "token_cache"]
    )
    assert client.post("/api/outlook/disconnect").status_code == 200
    with sessions() as db:
        connection = db.get(OutlookConnection, 1)
        assert connection.token_cache == "" and connection.delta_link is None
    assert client.post("/api/sync").status_code == 409


def test_expired_connection_flow_cannot_be_completed(microsoft, identity):
    sessions, _ = microsoft
    with sessions() as db:
        _, cookie = outlook_auth.begin_connection(db, 1, "http://localhost:8000")
        flow = db.scalar(select(OutlookAuthFlow))
        flow.expires_at = utcnow() - timedelta(seconds=1)
        db.commit()
        with pytest.raises(AppError):
            outlook_auth.complete_connection(db, cookie, {"state": "random-test-state", "code": "test-auth-code"})
        assert db.get(OutlookConnection, 1) is None


def test_callback_is_exempt_from_origin_filter_but_still_requires_browser_state(microsoft, identity):
    _, client = microsoft
    response = client.post(
        "/api/outlook/callback",
        data={"state": "forged", "code": "test-auth-code"},
        headers={"Origin": "https://login.microsoftonline.com", "Sec-Fetch-Site": "cross-site"},
        follow_redirects=False,
    )
    assert response.status_code == 303 and response.headers["location"] == "/?outlook=failed"
    assert client.post("/api/outlook/disconnect", headers={"Origin": "https://outside.invalid"}).status_code == 403


@pytest.mark.parametrize("authority", [TENANT, "common"])
def test_actual_msal_generates_pkce_and_nonce_without_network(microsoft, monkeypatch, authority):
    sessions, _ = microsoft
    monkeypatch.setattr(settings, "outlook_tenant_id", authority)

    class MetadataResponse:
        status_code = 200
        headers = {}

        @property
        def text(self):
            return json.dumps(self.json())

        def raise_for_status(self):
            pass

        def json(self):
            root = "https://login.microsoftonline.com/" + authority
            return {
                "authorization_endpoint": root + "/oauth2/v2.0/authorize",
                "token_endpoint": root + "/oauth2/v2.0/token",
                "issuer": root + "/v2.0",
                "jwks_uri": root + "/discovery/v2.0/keys",
            }

    class IdentityMetadata:
        trust_env = False

        def get(self, *args, **kwargs):
            return MetadataResponse()

        def post(self, *args, **kwargs):
            raise AssertionError("Sign-in initialization must not exchange a token")

    monkeypatch.setattr(outlook_auth.requests, "Session", IdentityMetadata)
    with sessions() as db:
        url, _ = outlook_auth.begin_connection(db, 1, "http://localhost:8000")
        query = parse_qs(urlsplit(url).query)
        assert query["code_challenge_method"] == ["S256"]
        assert query["nonce"] and query["state"]
        assert {"Mail.Read", "User.Read", "offline_access"} <= set(query["scope"][0].split())
        assert not ({"Mail.Send", "Mail.ReadWrite"} & set(query["scope"][0].split()))


def test_callback_access_logs_do_not_include_authorization_code():
    record = logging.LogRecord(
        "uvicorn.access",
        logging.INFO,
        "",
        0,
        '%s - "%s %s HTTP/%s" %d',
        ("client", "GET", "/api/outlook/callback?code=private-code&state=private-state", "1.1", 303),
        None,
    )
    OAuthAccessLogFilter().filter(record)
    assert "private-code" not in record.getMessage() and "private-state" not in record.getMessage()


def test_unconfigured_real_sync_does_not_fall_back_to_demo(microsoft, monkeypatch):
    sessions, client = microsoft
    monkeypatch.setattr(settings, "outlook_client_id", "")
    assert client.get("/api/outlook").json()["configured"] is False
    assert client.post("/api/outlook/connect").status_code == 409
    assert client.post("/api/sync").status_code == 409
    assert client.post("/api/demo").status_code == 409
    with sessions() as db:
        assert not mailbox.ready(db)


def mime_message(subject="[DCC] [2026-11] [WEEK:2026-10-31] Department summary"):
    mail = EmailMessage()
    mail["From"] = "department@hospital.example"
    mail["To"] = "hr@hospital.example"
    mail["Message-ID"] = "<real-sample@hospital.example>"
    mail["Subject"] = subject
    mail.set_content("Local integration test")
    content = workbook(
        [
            [
                "REAL-001",
                "EMP001",
                "Example Employee",
                "Cardiology",
                "2026-10-30",
                "AM",
                PUNCH_IN,
                message("tests.explanation.forgot_check"),
                "",
            ]
        ]
    )
    mail.add_attachment(
        content,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="summary.xlsx",
    )
    return mail.as_bytes()


def connected(db):
    connection = OutlookConnection(
        id=1,
        owner_id=1,
        status="CONNECTED",
        mailbox="hr@hospital.example",
        mailbox_id="mailbox-1",
        account_id="test-account",
        token_cache=outlook_auth.encrypt("{}"),
    )
    db.add(connection)
    db.commit()
    return connection


def configure_real_catalog(client):
    result = client.put("/api/departments/1", json={"name": "Cardiology", "senders": ["department@hospital.example"]})
    assert result.status_code == 200
    result = client.post(
        "/api/employees/import",
        content="employee_id,name,department_id\nEMP001,Example Employee,1\n",
        headers={"Content-Type": "text/csv"},
    )
    assert result.status_code == 200 and result.json()["imported"] == 1


def test_graph_delta_imports_real_workbook_and_separates_it_from_demo(microsoft, monkeypatch):
    sessions, client = microsoft
    with sessions() as db:
        ingest(db, workbook([row(code="LEGACY-DEMO")]))
    configure_real_catalog(client)
    seen = []
    checkpoint = "https://graph.microsoft.com/v1.0/me/mailFolders/inbox/messages/delta?$deltatoken=next"

    def handler(request):
        seen.append(request)
        assert request.method == "GET" and request.headers["Authorization"] == "Bearer private-access-token"
        if request.url.path.endswith("/delta"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {"id": "normal-mail", "subject": "Unrelated personal email"},
                        {
                            "id": "immutable-1",
                            "subject": "[DCC] Department summary",
                            "receivedDateTime": "2026-11-02T02:00:00Z",
                        },
                    ],
                    "@odata.deltaLink": checkpoint,
                },
            )
        assert request.url.path.endswith("/immutable-1/$value")
        return httpx.Response(200, content=mime_message())

    graph = httpx.Client(transport=httpx.MockTransport(handler))
    monkeypatch.setattr(outlook_auth, "access_token", lambda connection: "private-access-token")
    factory = OutlookProvider
    monkeypatch.setattr(mailbox, "OutlookProvider", lambda connection, token: factory(connection, token, client=graph))
    with sessions() as db:
        connection = connected(db)
        result = mailbox.synchronize_mailbox(db)
        db.commit()
        assert result["new"] == 1 and not result["has_more"]
        assert connection.delta_link == checkpoint
        email = db.scalar(select(Email).where(Email.source_id.startswith("outlook:")))
        process_email(db, email.id)
        db.commit()
        assert email.status == "IMPORTED" and email.period_id == "2026-11"
        assert mailbox.synchronize_mailbox(db)["existing"] == 1
        db.commit()
    assert len([request for request in seen if "/normal-mail/" in request.url.path]) == 0
    assert all("isRead" not in str(request.url) for request in seen)
    for source, expected in [("outlook", "REAL-001"), ("mailpit", "LEGACY-DEMO")]:
        scope = {"source": source}
        listing = client.get("/api/cases", params=scope).json()
        assert [case["code"] for case in listing["items"]] == [expected]
        assert len(client.get("/api/emails", params=scope).json()) == 1
        assert client.get("/api/dashboard", params=scope).json()["total"] == 1
        book = load_workbook(BytesIO(client.get("/api/export", params=scope).content), read_only=True)
        assert book.active.max_row == 2 and book.active.cell(2, 1).value == expected
    graph.close()


@pytest.mark.parametrize(
    "url",
    [
        "https://outside.invalid/v1.0/me/messages",
        "http://graph.microsoft.com/v1.0/me/messages",
        "https://graph.microsoft.com@outside.invalid/v1.0/",
        "https://graph.microsoft.com:444/v1.0/",
    ],
)
def test_graph_pagination_cannot_send_bearer_tokens_to_another_origin(url):
    with pytest.raises(ValueError):
        graph_url(url)


def test_cursor_is_not_advanced_when_a_page_fails(microsoft, monkeypatch):
    sessions, client = microsoft
    configure_real_catalog(client)
    checkpoint = "https://graph.microsoft.com/v1.0/me/mailFolders/inbox/messages/delta?$deltatoken=before"

    def handler(request):
        if request.url.path.endswith("/delta"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {"id": index, "subject": "[DCC] summary", "receivedDateTime": "2026-11-02T02:00:00Z"}
                        for index in ("one", "two")
                    ],
                    "@odata.deltaLink": checkpoint + "&after=1",
                },
            )
        return httpx.Response(200, content=mime_message()) if "/one/" in request.url.path else httpx.Response(503)

    graph = httpx.Client(transport=httpx.MockTransport(handler))
    factory = OutlookProvider
    monkeypatch.setattr(outlook_auth, "access_token", lambda connection: "private-access-token")
    monkeypatch.setattr(mailbox, "OutlookProvider", lambda connection, token: factory(connection, token, client=graph))
    with sessions() as db:
        connection = connected(db)
        connection.delta_link = checkpoint
        db.commit()
        with pytest.raises(AppError):
            mailbox.synchronize_mailbox(db)
        db.rollback()
    with sessions() as db:
        assert db.get(OutlookConnection, 1).delta_link == checkpoint
        assert db.scalar(select(Email)) is None
    graph.close()


def test_employee_master_import_is_atomic_and_does_not_overwrite_existing_people(microsoft):
    _, client = microsoft
    content = "employee_id,name,department_id\nNEW-1,Example One,1\nNEW-2,Example Two,99999\n"
    assert client.post("/api/employees/import", content=content).status_code == 422
    valid = "employee_id,name,department_id\nNEW-1,Example One,1\n"
    assert client.post("/api/employees/import", content=valid).json()["imported"] == 1
    assert client.post("/api/employees/import", content=valid).json()["imported"] == 0
    conflict = valid.replace("Example One", "Different Person")
    assert client.post("/api/employees/import", content=conflict).status_code == 409


@pytest.mark.parametrize("scheme", ["http", "https"])
def test_callback_route_completes_only_the_browser_started_flow(microsoft, identity, monkeypatch, scheme):
    _, client = microsoft
    origin = scheme + "://localhost:8000"
    monkeypatch.setattr(settings, "outlook_redirect_uri", origin + "/api/outlook/callback")
    client.post(
        origin + "/api/auth/login", json={"username": "hr.demo", "password": "Test-demo-password"}
    ).raise_for_status()
    start = client.post(origin + "/api/outlook/connect")
    assert start.status_code == 200
    cookie = start.headers["set-cookie"].lower()
    assert "httponly" in cookie and "path=/api/outlook/callback" in cookie
    assert ("samesite=none" in cookie and "secure" in cookie) if scheme == "https" else "samesite=lax" in cookie
    values = {"state": "random-test-state", "code": "test-auth-code"}
    if scheme == "https":
        result = client.post(
            origin + "/api/outlook/callback",
            data=values,
            headers={"Origin": "https://login.microsoftonline.com", "Sec-Fetch-Site": "cross-site"},
            follow_redirects=False,
        )
    else:
        result = client.get(origin + "/api/outlook/callback", params=values, follow_redirects=False)
    assert result.status_code == 303 and result.headers["location"] == "/?outlook=connected"
    assert client.get(origin + "/api/outlook").json()["status"] == "CONNECTED"


def test_department_mail_without_required_subject_infers_period_from_file(microsoft, monkeypatch):
    from app.services.imports import synchronize
    from app.services.outlook_provider import OutlookProvider

    sessions, client = microsoft
    configure_real_catalog(client)
    mail = EmailMessage()
    mail["From"] = "department@hospital.example"
    mail["To"] = "team-distribution@hospital.example"
    mail["Message-ID"] = "<ordinary-department-mail@hospital.example>"
    mail["Subject"] = "Monthly attendance corrections"
    mail.set_content("Department report")
    rows = [
        [
            "ORDINARY-001",
            "EMP001",
            "Example Employee",
            "Cardiology",
            "2026-09-30",
            "AM",
            PUNCH_IN,
            message("tests.explanation.forgot_check"),
            "",
        ]
    ]
    mail.add_attachment(
        workbook(rows),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="summary.xlsx",
    )
    end = "https://graph.microsoft.com/v1.0/me/mailFolders/inbox/messages/delta?$deltatoken=end"

    def handler(request):
        if request.url.path.endswith("/delta"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "ordinary",
                            "subject": "Monthly attendance corrections",
                            "from": {"emailAddress": {"address": "department@hospital.example"}},
                            "receivedDateTime": "2026-11-02T02:00:00Z",
                        }
                    ],
                    "@odata.deltaLink": end,
                },
            )
        return httpx.Response(200, content=mail.as_bytes())

    with httpx.Client(transport=httpx.MockTransport(handler)) as graph, sessions() as db:
        connection = connected(db)
        provider = OutlookProvider(connection, "test-token", client=graph)
        provider.allowed_senders = {"department@hospital.example"}
        assert synchronize(db, provider)["new"] == 1
        email = db.scalar(select(Email))
        assert email.period_id == "2026-10" and email.cycle_start is None
        process_email(db, email.id)
        db.commit()
        assert email.status == "IMPORTED"
        assert db.scalar(select(AttendanceCase)).period_id == "2026-10"


def test_delta_pagination_and_expired_cursor_resume_without_losing_checkpoint(microsoft):
    sessions, _ = microsoft
    calls = []
    next_page = "https://graph.microsoft.com/v1.0/me/mailFolders/inbox/messages/delta?$skiptoken=page2"
    final = "https://graph.microsoft.com/v1.0/me/mailFolders/inbox/messages/delta?$deltatoken=final"

    def handler(request):
        calls.append(str(request.url))
        if "expired" in str(request.url):
            return httpx.Response(410)
        if "page2" in str(request.url):
            return httpx.Response(200, json={"value": [], "@odata.deltaLink": final})
        return httpx.Response(200, json={"value": [], "@odata.nextLink": next_page})

    with httpx.Client(transport=httpx.MockTransport(handler)) as graph, sessions() as db:
        connection = connected(db)
        connection.delta_link = final + "&expired=1"
        first = OutlookProvider(connection, "test-token", client=graph)
        assert list(first.messages()) == [] and first.has_more and first.checkpoint == next_page
        connection.delta_link = first.checkpoint
        second = OutlookProvider(connection, "test-token", client=graph)
        assert list(second.messages()) == [] and not second.has_more and second.checkpoint == final
        assert len(calls) == 3


@pytest.mark.parametrize("status,expected", [(403, 409), (429, 503)])
def test_graph_permission_and_throttle_errors_are_actionable(microsoft, monkeypatch, status, expected):
    sessions, _ = microsoft
    factory = OutlookProvider
    with httpx.Client(
        transport=httpx.MockTransport(lambda request: httpx.Response(status, headers={"Retry-After": "120"}))
    ) as graph:
        monkeypatch.setattr(outlook_auth, "access_token", lambda connection: "test-token")
        monkeypatch.setattr(
            mailbox, "OutlookProvider", lambda connection, token: factory(connection, token, client=graph)
        )
        with sessions() as db:
            connection = connected(db)
            with pytest.raises(AppError) as captured:
                mailbox.synchronize_mailbox(db)
            assert captured.value.status_code == expected
            assert connection.delta_link is None
            if status == 429:
                assert captured.value.retry_after == 120


def test_department_creation_handles_duplicate_addresses_without_server_error(microsoft):
    _, client = microsoft
    department = {"name": "New Ward", "senders": ["new-ward@hospital.example"]}
    created = client.post("/api/departments", json=department)
    assert created.status_code == 200
    assert created.json()["id"] > 0
    assert client.post("/api/departments", json=department).status_code == 409
    assert client.put("/api/departments/2", json=department).status_code == 409
    assert (
        client.post("/api/departments", json={"name": "   ", "senders": ["other@hospital.example"]}).status_code == 422
    )


@pytest.mark.parametrize("tenant", [TENANT, "9188040d-6c67-4c5b-b112-36a304b66dad"])
def test_common_authority_accepts_the_account_chosen_in_microsoft(microsoft, monkeypatch, tenant):
    sessions, client = microsoft
    monkeypatch.setattr(settings, "outlook_tenant_id", "common")
    monkeypatch.setattr(settings, "outlook_test_mode", True)

    class TestIdentity(FakeIdentity):
        def initiate_auth_code_flow(self, scopes, **kwargs):
            assert not kwargs.get("login_hint") and kwargs["prompt"] == "select_account"
            return super().initiate_auth_code_flow(scopes, **kwargs)

        def acquire_token_by_auth_code_flow(self, flow, parameters):
            result = super().acquire_token_by_auth_code_flow(flow, parameters)
            result["id_token_claims"]["tid"] = tenant
            return result

    monkeypatch.setattr(outlook_auth, "microsoft_client", lambda cache: TestIdentity(cache))
    monkeypatch.setattr(
        outlook_auth, "read_profile", lambda token: {"id": "mailbox-test", "mail": "Developer@Example.edu"}
    )
    with sessions() as db:
        _, cookie = outlook_auth.begin_connection(db, 1, "http://localhost:8000")
        outlook_auth.complete_connection(db, cookie, {"state": "random-test-state", "code": "test-auth-code"})
        assert db.get(OutlookConnection, 1).mailbox == "developer@example.edu"
    status = client.get("/api/outlook").json()
    assert status["test_mode"] is True and "allowed_mailbox" not in status


@pytest.mark.parametrize("address", ["first@example.edu", "second@outlook.com"])
def test_account_selection_queues_initial_sync_without_an_email_allowlist(microsoft, identity, monkeypatch, address):
    sessions, _ = microsoft
    monkeypatch.setattr(settings, "outlook_test_mode", True)
    monkeypatch.setattr(outlook_auth, "read_profile", lambda token: {"id": "selected-mailbox", "mail": address})
    with sessions() as db:
        _, cookie = outlook_auth.begin_connection(db, 1, "http://localhost:8000")
        outlook_auth.complete_connection(db, cookie, {"state": "random-test-state", "code": "test-auth-code"})
        assert db.get(OutlookConnection, 1).mailbox == address
        job = db.scalar(select(ProcessingJob).where(ProcessingJob.kind == "SYNC"))
        assert job.status == "PENDING" and "token" not in str(job.payload)


def test_refresh_uses_the_connected_microsoft_account(microsoft, monkeypatch):
    sessions, _ = microsoft

    class RefreshIdentity(FakeIdentity):
        def get_accounts(self):
            return [{"home_account_id": "another-account"}, {"home_account_id": "test-account"}]

        def acquire_token_silent_with_error(self, scopes, account):
            assert account["home_account_id"] == "test-account"
            return {"access_token": "selected-access-token"}

    monkeypatch.setattr(outlook_auth, "microsoft_client", lambda cache: RefreshIdentity(cache))
    with sessions() as db:
        assert outlook_auth.access_token(connected(db)) == "selected-access-token"


def test_mvp_downloads_only_tagged_mail_even_from_a_configured_sender(microsoft, monkeypatch):
    sessions, client = microsoft
    configure_real_catalog(client)
    monkeypatch.setattr(settings, "outlook_test_mode", True)
    subject = "[TALLYV-MVP] [DCC] [2026-11] [WEEK:2026-10-31] Test summary"
    checkpoint = "https://graph.microsoft.com/v1.0/me/mailFolders/inbox/messages/delta?$deltatoken=test"
    downloaded = []

    def handler(request):
        if request.url.path.endswith("/delta"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": key,
                            "subject": title,
                            "from": {"emailAddress": {"address": "department@hospital.example"}},
                            "receivedDateTime": "2026-11-02T02:00:00Z",
                        }
                        for key, title in [
                            ("personal", "Personal email"),
                            ("business", "[DCC] Department report"),
                            ("incomplete", "[TALLYV-MVP] Unrelated test"),
                            ("test", subject),
                        ]
                    ],
                    "@odata.deltaLink": checkpoint,
                },
            )
        downloaded.append(request.url.path)
        return httpx.Response(200, content=mime_message(subject))

    with httpx.Client(transport=httpx.MockTransport(handler)) as graph, sessions() as db:
        connection = connected(db)
        provider = OutlookProvider(connection, "test-token", client=graph)
        provider.allowed_senders = {"department@hospital.example"}
        from app.services.imports import synchronize

        assert synchronize(db, provider)["new"] == 1
        email = db.scalar(select(Email))
        assert email.cycle_start.isoformat() == "2026-10-31" and email.period_id == "2026-11"
        process_email(db, email.id)
        db.commit()
        assert email.status == "IMPORTED" and provider.checkpoint == checkpoint
        assert db.scalar(select(AttendanceCase)).code == "REAL-001"
    assert downloaded == ["/v1.0/me/messages/test/$value"]


def test_fixed_tenant_still_rejects_other_organizations(microsoft, identity, monkeypatch):
    sessions, _ = microsoft
    monkeypatch.setattr(settings, "outlook_tenant_id", "33333333-3333-4333-8333-333333333333")
    with sessions() as db:
        _, cookie = outlook_auth.begin_connection(db, 1, "http://localhost:8000")
        with pytest.raises(AppError) as error:
            outlook_auth.complete_connection(db, cookie, {"state": "random-test-state", "code": "test-auth-code"})
        assert error.value.status_code == 403 and db.get(OutlookConnection, 1) is None


def test_multiple_handlers_are_saved_and_grouped_into_the_same_department(microsoft):
    from types import SimpleNamespace

    from app.services.email_provider import IncomingEmail
    from app.services.imports import synchronize

    sessions, client = microsoft
    configure_real_catalog(client)
    response = client.put(
        "/api/departments/1",
        json={
            "name": "Cardiology",
            "senders": ["department@hospital.example", " SECOND@hospital.example ", "second@hospital.example"],
        },
    )
    assert response.status_code == 200
    assert response.json()["senders"] == ["department@hospital.example", "second@hospital.example"]
    template = mime_message()
    from email import policy
    from email.parser import BytesParser

    parsed = BytesParser(policy=policy.default).parsebytes(template)
    attachment = next(parsed.iter_attachments()).get_payload(decode=True)
    incoming = [
        IncomingEmail(
            "outlook:handler-" + str(index),
            "<handler-" + str(index) + "@example.test>",
            sender,
            "hr@hospital.example",
            "[DCC] [2026-11] [WEEK:2026-10-31] Test",
            "Test summary",
            utcnow(),
            [("summary.xlsx", attachment)],
        )
        for index, sender in enumerate(response.json()["senders"])
    ]
    provider = SimpleNamespace(mailbox_address="hr@hospital.example", messages=lambda: iter(incoming))
    with sessions() as db:
        assert synchronize(db, provider)["new"] == 2
        for email in list(db.scalars(select(Email))):
            assert email.department_id == 1
            process_email(db, email.id)
        db.commit()
        assert len(list(db.scalars(select(AttendanceCase)))) == 1
    assert client.get("/api/dashboard", params={"source": "outlook"}).json()["departments_sent"] == 1
    assert (
        client.put(
            "/api/departments/1", json={"name": "Cardiology", "senders": ["second@hospital.example"]}
        ).status_code
        == 200
    )
    with sessions() as db:
        from app.crud.catalog import find_department_by_sender

        assert find_department_by_sender(db, "department@hospital.example") is None
        assert find_department_by_sender(db, "SECOND@HOSPITAL.EXAMPLE").id == 1


@pytest.mark.parametrize("department_name,expected", [("Cardiology", 1), ("Orthopedics", 2), ("Unmapped", None)])
def test_handler_shared_by_departments_is_routed_using_the_workbook(microsoft, department_name, expected):
    from types import SimpleNamespace

    from app.services.email_provider import IncomingEmail
    from app.services.imports import synchronize

    sessions, client = microsoft
    for identifier, name in [(1, "Cardiology"), (2, "Orthopedics")]:
        assert (
            client.put(
                "/api/departments/" + str(identifier), json={"name": name, "senders": ["shared@example.test"]}
            ).status_code
            == 200
        )
    content = workbook(
        [
            [
                "SHARED-001",
                "EXAMPLE-001",
                "Example Person",
                department_name,
                "2026-10-30",
                "AM",
                PUNCH_IN,
                "Test explanation",
                "",
            ]
        ]
    )
    incoming = IncomingEmail(
        "outlook:shared-test",
        "<shared@example.test>",
        "shared@example.test",
        "hr@hospital.example",
        "[DCC] [2026-11] [WEEK:2026-10-31] Test",
        "Test",
        utcnow(),
        [("summary.xlsx", content)],
    )
    provider = SimpleNamespace(mailbox_address="hr@hospital.example", messages=lambda: iter([incoming]))
    with sessions() as db:
        assert synchronize(db, provider)["new"] == 1
        email = db.scalar(select(Email))
        assert email.department_id == expected
        assert email.status == ("RECEIVED" if expected else "NEEDS_REVIEW")


def test_sender_configuration_validates_addresses_and_supports_no_handlers_yet(microsoft):
    _, client = microsoft
    assert client.post("/api/departments", json={"name": "Draft Department", "senders": []}).json()["senders"] == []
    assert (
        client.post("/api/departments", json={"name": "Invalid Department", "senders": ["invalid"]}).status_code == 422
    )
    assert (
        client.post("/api/departments", json={"name": "Legacy Department", "sender": "one@example.test"}).status_code
        == 422
    )
