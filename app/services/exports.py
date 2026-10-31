from datetime import timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.core.attendance import EXPORT_HEADERS, EXPORT_SHEET_NAME
from app.crud import cases as case_crud
from app.schema.cases import CaseFilters


def create_workbook(db: Session, selection: CaseFilters) -> BytesIO:
    rows = case_crud.list_rows(db, selection, ordered=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = EXPORT_SHEET_NAME
    headers = EXPORT_HEADERS
    sheet.append(headers)
    for case, version, analysis, department, email in rows:
        n = version.normalized
        values = [
            case.code,
            case.employee_id,
            n["name"],
            department.name,
            n["date"],
            n["shift"],
            n["punch"],
            n["explanation"],
            analysis.category if analysis.status == "COMPLETED" else "",
            analysis.assessment if analysis.status == "COMPLETED" else "NOT_EVALUATED",
            case.hr_category or "",
            case.hr_status,
            case.hr_note,
            case.confirmed_by or "",
            case.confirmed_at.isoformat() if case.confirmed_at else "",
            case.period_id,
            email.message_id or email.source_id,
            str(version.attachment_id),
            version.sheet,
            str(version.row_number),
            case.cycle_start.isoformat() if case.cycle_start else "",
            (case.cycle_start + timedelta(days=2)).isoformat() if case.cycle_start else "",
        ]
        row_index = sheet.max_row + 1
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row_index, col, str(value))
            cell.data_type = "s"  # Force literal strings, including =, +, -, @, tabs.
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="176B5B")
        cell.font = Font(color="FFFFFF", bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(48, max(16, len(str(column[0].value)) + 3))
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
