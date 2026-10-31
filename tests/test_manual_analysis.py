from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import func, select

from app import worker
from app.core.i18n import message
from app.demo.generate import picture, workbook
from app.models import AnalysisResult, AttendanceCase, CaseVersion, EvidenceFile, ProcessingJob
from app.services.imports import synchronize
from tests.helpers import Provider, ingest, row


def test_scan_imports_before_explicit_analysis_and_repeated_click_does_not_queue_twice(context, monkeypatch):
    sessions, client = context
    monkeypatch.setattr(worker, "SessionLocal", sessions)
    with sessions() as db:
        synchronize(db, Provider(workbook([row()], {2: [picture("MP001", "2026-10-01")]})))
        db.commit()
    assert client.get("/api/analysis").json()["importing"] is True
    assert client.post("/api/analysis").status_code == 409
    assert worker.run_one()  # Import Excel; do not schedule OCR/classification.
    assert not worker.run_one()
    with sessions() as db:
        assert db.scalar(select(func.count()).select_from(AttendanceCase)) == 1
        result = db.scalar(select(AnalysisResult))
        assert (result.status, result.assessment) == ("NOT_STARTED", "NOT_EVALUATED")
        assert db.scalar(select(EvidenceFile)).ocr_text == ""
        assert db.scalar(select(ProcessingJob).where(ProcessingJob.kind == "ANALYZE")) is None
    case = client.get("/api/cases").json()["items"][0]
    assert not case["bulk_eligible"]
    assert client.get("/api/cases?category=UNKNOWN").json()["total"] == 0
    exported = load_workbook(BytesIO(client.get("/api/export").content))
    assert exported.active.cell(2, 9).value is None
    assert exported.active.cell(2, 10).value == "NOT_EVALUATED"
    assert client.get("/api/analysis").json() == {
        "waiting": 1,
        "processing": 0,
        "completed": 0,
        "failed": 0,
        "importing": False,
    }
    assert client.post("/api/analysis").json() == {"queued": 1}
    assert client.post("/api/analysis").json() == {"queued": 0}
    assert client.get("/api/analysis").json()["processing"] == 1
    assert worker.run_one()
    with sessions() as db:
        result = db.scalar(select(AnalysisResult))
        assert result.status == "COMPLETED" and result.category == "FORGOT_CHECK"
        assert db.scalar(select(AttendanceCase)).hr_status == "PENDING_REVIEW"
    assert client.post("/api/analysis").json() == {"queued": 0}
    assert client.get("/api/analysis").json()["completed"] == 1


def test_analysis_scope_includes_proposals_excludes_other_sources_and_preserves_completed(context):
    sessions, client = context
    with sessions() as db:
        ingest(db, workbook([row()]))
        first = db.scalar(select(AnalysisResult))
        first.status, first.category, first.assessment = "COMPLETED", "FORGOT_CHECK", "SUFFICIENT"
        db.commit()
        ingest(db, workbook([row(explanation=message("tests.explanation.revised_device"))]), "source-2")
        ingest(db, workbook([row("OUTLOOK-001")]), "outlook:test:other")
        proposed = db.scalar(select(CaseVersion).where(CaseVersion.status == "PROPOSED"))
        proposed_id, first_id = proposed.id, first.id
    assert client.post("/api/analysis?source=mailpit&period=2026-09").json() == {"queued": 0}
    assert client.post("/api/analysis?source=mailpit&period=2026-10&department=2").json() == {"queued": 0}
    assert client.post("/api/analysis?source=mailpit&period=2026-10&department=1").json() == {"queued": 1}
    with sessions() as db:
        assert db.get(AnalysisResult, first_id).status == "COMPLETED"
        assert db.scalar(select(AnalysisResult).where(AnalysisResult.version_id == proposed_id)).status == "PENDING"
        waiting = list(db.scalars(select(AnalysisResult).where(AnalysisResult.status == "NOT_STARTED")))
        assert len(waiting) == 1
        assert db.scalar(select(func.count()).select_from(ProcessingJob).where(ProcessingJob.kind == "ANALYZE")) == 1


def test_failed_analysis_can_be_requested_again_and_endpoints_require_authentication(context):
    sessions, client = context
    with sessions() as db:
        ingest(db, workbook([row()]))
        result = db.scalar(select(AnalysisResult))
        result.status, result.error = "FAILED", "A previous error"
        case = db.scalar(select(AttendanceCase))
        case.hr_status, case.confirmed_by = "CONFIRMED", "hr.demo"
        db.commit()
    assert client.get("/api/analysis").json()["failed"] == 1
    assert client.post("/api/analysis").json() == {"queued": 1}
    with sessions() as db:
        result = db.scalar(select(AnalysisResult))
        assert result.status == "PENDING" and result.error is None
        case = db.scalar(select(AttendanceCase))
        assert case.hr_status == "PENDING_REVIEW" and case.confirmed_by is None
    client.post("/api/auth/logout")
    assert client.get("/api/analysis").status_code == 401
    assert client.post("/api/analysis").status_code == 401


def test_worker_failure_is_visible_and_manual_retry_recovers(context, monkeypatch):
    sessions, client = context
    monkeypatch.setattr(worker, "SessionLocal", sessions)
    with sessions() as db:
        synchronize(db, Provider(workbook([row()])))
        db.commit()
    assert worker.run_one()
    assert client.post("/api/analysis").json() == {"queued": 1}
    with sessions() as db:
        job = db.scalar(select(ProcessingJob).where(ProcessingJob.kind == "ANALYZE"))
        job.max_attempts = 1
        db.commit()
    actual = worker.analyze

    def fail(db, version_id):
        raise RuntimeError("simulated analysis failure")

    monkeypatch.setattr(worker, "analyze", fail)
    assert worker.run_one()
    assert client.get("/api/analysis").json()["failed"] == 1
    assert client.get("/api/analysis").json()["processing"] == 0
    monkeypatch.setattr(worker, "analyze", actual)
    assert client.post("/api/analysis").json() == {"queued": 1}
    assert worker.run_one()
    assert client.get("/api/analysis").json()["completed"] == 1
