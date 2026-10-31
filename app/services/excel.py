"""Parse untrusted Excel rows and image anchors without database or network access."""

import hashlib
import re
from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor
from PIL import Image

from app.core.attendance import HEADERS, PUNCH_IN, PUNCH_NA, PUNCH_OUT, SHEET_NAME
from app.core.i18n import message

KEYS = ["code", "employee_id", "name", "department", "date", "shift", "punch", "explanation", "evidence"]
MAX_BYTES = 15 * 1024 * 1024


def strict_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    # ISO is deliberately the only string format. 01/02/2026 is ambiguous.
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        raise ValueError(message("excel.ambiguous_date"))
    return date.fromisoformat(text)


def xlsx_guard(content):
    if len(content) > MAX_BYTES:
        raise ValueError(message("excel.file_too_large"))
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(x.file_size for x in members) > 60 * 1024 * 1024:
                raise ValueError(message("excel.archive_too_large"))
            if any("vbaProject" in x.filename or "externalLinks/" in x.filename for x in members):
                raise ValueError(message("excel.external_content_denied"))
    except BadZipFile:
        raise ValueError(message("excel.invalid_workbook"))


def parse_workbook(content):
    xlsx_guard(content)
    wb = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(message("excel.missing_sheet"))
    ws = wb[SHEET_NAME]
    if ws.max_row > 2001 or ws.max_column > 30:
        raise ValueError(message("excel.sheet_limits_exceeded"))
    rows = {
        r: [ws.cell(r, c) for c in range(1, 10)]
        for r in range(2, ws.max_row + 1)
        if any(ws.cell(r, c).value is not None for c in range(1, 10))
    }
    actual = [str(ws.cell(1, c).value or "").strip() for c in range(1, 10)]
    if actual != HEADERS:
        return (
            [],
            [
                {
                    "row": r,
                    "codes": [message("excel.required_columns_missing")],
                    "raw": {KEYS[i]: str(cell.value or "") for i, cell in enumerate(cells)},
                }
                for r, cells in rows.items()
            ],
            len(rows),
            [message("excel.template_mismatch")],
        )
    images = {}
    bad_rows, global_errors = set(), []
    for index, picture in enumerate(ws._images):
        anchor = picture.anchor
        if not isinstance(anchor, (OneCellAnchor, TwoCellAnchor)):
            global_errors.append(message("excel.anchor_row_unknown", value0=index + 1))
            continue
        start = anchor._from
        row = start.row + 1
        ambiguous = isinstance(anchor, TwoCellAnchor) and anchor.to.row != start.row
        if start.col != 8 or row not in rows or start.rowOff < 0 or ambiguous:
            global_errors.append(message("excel.anchor_ambiguous", value0=index + 1, value1=row))
            if row in rows:
                bad_rows.add(row)
            continue
        try:
            image_content = picture._data()
            if len(image_content) > 8 * 1024 * 1024:
                raise ValueError(message("excel.image_too_large"))
            with Image.open(BytesIO(image_content)) as image:
                if image.width * image.height > 16_000_000:
                    raise ValueError(message("excel.image_pixel_limit"))
                image.load()
                output = BytesIO()
                image.convert("RGB").save(output, format="PNG")
                image_content = output.getvalue()
            images.setdefault(row, []).append(
                {
                    "content": image_content,
                    "hash": hashlib.sha256(image_content).hexdigest(),
                    "anchor": f"I{row}#{index + 1}",
                }
            )
        except Exception:
            bad_rows.add(row)
            global_errors.append(message("excel.image_unreadable", value0=index + 1))
    parsed, errors = [], []
    for row, cells in rows.items():
        raw = {
            key: (cell.value.isoformat() if isinstance(cell.value, (date, datetime)) else str(cell.value or ""))
            for key, cell in zip(KEYS, cells)
        }
        issues = []
        if any(cell.data_type == "f" for cell in cells):
            issues.append(message("excel.formula_denied"))
        if any(len(value) > 10000 for value in raw.values()):
            issues.append(message("excel.cell_too_long"))
        if row in bad_rows:
            issues.append(message("excel.evidence_link_invalid"))
        normalized = {key: value.strip() for key, value in raw.items()}
        for key in ["code", "employee_id", "name", "department", "date"]:
            if not normalized[key]:
                issues.append(message("excel.required_field_missing", value0=key))
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,60}", normalized["code"]):
            issues.append(message("excel.case_code_invalid"))
        try:
            normalized["date"] = strict_date(cells[4].value).isoformat()
        except ValueError as error:
            issues.append(str(error))
        if normalized["punch"] not in {"", PUNCH_IN, PUNCH_OUT, PUNCH_NA}:
            issues.append(message("excel.punch_type_invalid"))
        if issues:
            errors.append({"row": row, "codes": issues, "raw": raw})
        else:
            parsed.append({"row": row, "raw": raw, "normalized": normalized, "images": images.get(row, [])})
    wb.close()
    return parsed, errors, len(rows), global_errors
