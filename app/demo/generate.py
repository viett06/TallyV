"""Generate ONLY synthetic workbooks/images, then send them via local SMTP.
No AttendanceCase records are inserted here. Answers live in a separate JSON file.
"""

import json
import smtplib
from datetime import date, datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont

from app.core.attendance import PUNCH_IN, PUNCH_NA, PUNCH_OUT, SHEET_NAME
from app.core.config import settings
from app.core.i18n import message
from app.demo.catalog import DEPARTMENTS
from app.services.excel import HEADERS
from app.services.review_calendar import business_today, default_cycle, parse_cycle, payroll_period, period_bounds

EXPLANATIONS = [
    message("demo.explanation.complete_forgot"),
    message("demo.explanation.complete_device"),
    message("demo.explanation.complete_leave"),
    "",
    message("demo.explanation.missing_image"),
    message("demo.explanation.unreadable_image"),
    message("demo.explanation.date_mismatch"),
    message("demo.explanation.employee_mismatch"),
    message("demo.explanation.ambiguous"),
    message("demo.explanation.out_of_scope"),
    message("demo.explanation.missing_shift"),
    message("demo.explanation.multiple_images"),
]
EXPECTED = [
    "FORGOT_CHECK",
    "DEVICE_ERROR",
    "LEAVE_NOT_RECORDED",
    "UNKNOWN",
    "FORGOT_CHECK",
    "DEVICE_ERROR",
    "FORGOT_CHECK",
    "DEVICE_ERROR",
    "UNKNOWN",
    "UNKNOWN",
    "FORGOT_CHECK",
    "FORGOT_CHECK",
]
GROUPS = [
    "complete_forgot",
    "complete_device",
    "complete_leave",
    "missing_explanation",
    "missing_image",
    "unreadable",
    "date_conflict",
    "employee_conflict",
    "ambiguous",
    "out_of_scope",
    "missing_shift_punch",
    "multiple_images",
]


def font(size):
    for candidate in [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default(size=size)


def picture(employee, day, kind="ATTENDANCE", unreadable=False):
    image = Image.new("RGB", (940, 410), "#f2f7f8")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 940, 65), fill="#155f52")
    draw.text((25, 15), message("demo.watermark"), font=font(28), fill="white")
    if unreadable:
        for x in range(30, 900, 10):
            draw.line((x, 110, x - 20, 300), fill="#bbbbbb", width=2)
        draw.text((25, 340), message("demo.unreadable_image_caption"), font=font(20), fill="#4c6660")
    elif kind == "CAMERA":
        draw.rectangle((40, 100, 900, 290), outline="#526d65", width=3)
        draw.text((70, 140), "CAMERA ILLUSTRATION - NO REAL PEOPLE", font=font(23), fill="#273f37")
        draw.text((70, 200), "Empty corridor / simulated scene", font=font(23), fill="#273f37")
    else:
        lines = [
            f"EMPLOYEE: {employee}",
            f"DATE: {day}",
            "TIME: 08:00:00",
            {"LEAVE": "STATUS: LEAVE APPROVED", "DEVICE": "STATUS: DEVICE ERROR", "ATTENDANCE": "STATUS: RECORDED"}[
                kind
            ],
        ]
        for idx, line in enumerate(lines):
            draw.text((28, 90 + idx * 57), line, font=font(32), fill="#192b25")
    output = BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def workbook(rows, pictures=None, bad_template=False, ambiguous_rows=None):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS[:-1] if bad_template else HEADERS)
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="176B5B")
    for row in rows:
        ws.append(row)
    for row_number, images in (pictures or {}).items():
        for image_index, content in enumerate(images):
            image = XLImage(BytesIO(content))
            image.width, image.height = 250, 110
            column = "H" if row_number in (ambiguous_rows or []) else "I"
            ws.add_image(image, f"{column}{row_number}")
        ws.row_dimensions[row_number].height = 100
    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 30
    ws.column_dimensions["H"].width = 70
    ws.freeze_panes = "C2"
    result = BytesIO()
    wb.save(result)
    return result.getvalue()


def send(smtp, index, department, period, subject, attachments, body=None, cycle_start=None):
    email = EmailMessage()
    email["From"] = f"department{department}@demo.invalid" if department else "unknown@demo.invalid"
    email["To"] = "hr@demo.invalid"
    week_tag = f" [WEEK:{cycle_start}]" if cycle_start else ""
    email["Subject"] = f"[DCC] [{period}]{week_tag} {subject}" if period else subject
    email["Message-ID"] = f"<tallyv-demo-v2-{cycle_start or period}-{index:02}@demo.invalid>"
    email["Date"] = format_datetime(datetime.now(timezone.utc))
    email.set_content(body or message("demo.email_body"))
    for name, content in attachments:
        email.add_attachment(
            content,
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=name,
        )
    smtp.send_message(email)


def run(period_id=None, cycle_start=None):
    today = business_today()
    period = period_id or payroll_period(today)
    cycle = parse_cycle(period, cycle_start) if cycle_start else default_cycle(period, today)
    starts, _ = period_bounds(period)
    first_day = starts.isoformat()
    if settings.smtp_host not in {"127.0.0.1", "localhost", "mailpit"}:
        raise ValueError("Demo SMTP must be the local Mailpit instance")
    demo_dir = settings.storage_dir / "demo"
    demo_dir.mkdir(parents=True, exist_ok=True)
    answers, files = {}, []
    with smtplib.SMTP(settings.smtp_host, settings.smtp_port, timeout=20) as smtp:
        for dep in range(1, 6):
            rows, images = [], {}
            for i in range(12):
                employee_index = (dep - 1) * 8 + i % 8 + 1
                employee = f"MP{employee_index:03}"
                code = f"MP-{cycle.isoformat()}-{dep}-{i + 1:03}"
                day = (starts + timedelta(days=i % ((cycle - starts).days + 1))).isoformat()
                shift = (
                    ""
                    if i == 10
                    else (message("demo.morning_shift") if i % 2 == 0 else message("demo.afternoon_shift"))
                )
                punch = "" if i == 10 else (PUNCH_NA if i == 2 else (PUNCH_IN if i % 2 == 0 else PUNCH_OUT))
                rows.append(
                    [
                        code,
                        employee,
                        message("demo.employee_name", value0=employee_index),
                        DEPARTMENTS[dep - 1],
                        day,
                        shift,
                        punch,
                        EXPLANATIONS[i],
                        "",
                    ]
                )
                kind = "LEAVE" if i == 2 else ("DEVICE" if i in {1, 5, 7} else "ATTENDANCE")
                image_day = (date.fromisoformat(day) + timedelta(days=1)).isoformat() if i == 6 else day
                image_employee = "MP999" if i == 7 else employee
                if i != 4:
                    images[i + 2] = [picture(image_employee, image_day, kind, unreadable=i == 5)]
                    if i == 11:
                        images[i + 2].append(picture(employee, day))
                    if i == 9:
                        images[i + 2] = [picture(employee, day, "CAMERA")]
                answers[code] = {"category": EXPECTED[i], "group": GROUPS[i]}
            # A valid file may contain bad rows: keep every row accounted for.
            rows.append(
                [
                    f"BAD-{dep}",
                    "MP999",
                    message("demo.unknown_employee"),
                    DEPARTMENTS[dep - 1],
                    first_day,
                    message("demo.short_morning_shift"),
                    PUNCH_IN,
                    message("category.forgot_check"),
                    "",
                ]
            )
            if dep == 1:
                rows.append(
                    [
                        "BAD-DATE",
                        "MP001",
                        message("demo.first_employee"),
                        DEPARTMENTS[0],
                        f"01/02/{starts.year:04}",
                        message("demo.short_morning_shift"),
                        PUNCH_IN,
                        message("category.forgot_check"),
                        "",
                    ]
                )
                # Deliberately anchor an image outside the evidence column.
                rows.append(
                    [
                        "BAD-ANCHOR",
                        "MP001",
                        message("demo.first_employee"),
                        DEPARTMENTS[0],
                        first_day,
                        message("demo.short_morning_shift"),
                        PUNCH_IN,
                        message("category.forgot_check"),
                        "",
                    ]
                )
                images[16] = [picture("MP001", first_day)]
            content = workbook(rows, images, ambiguous_rows={16} if dep == 1 else None)
            name = f"attendance-department-{dep}.xlsx"
            (demo_dir / name).write_bytes(content)
            files.append((name, content))
            send(
                smtp,
                dep,
                dep,
                period,
                message("demo.email.department_summary", value0=dep),
                [(name, content)],
                cycle_start=cycle,
            )
        send(smtp, 6, 1, period, message("demo.email.duplicate_file"), [files[0]], cycle_start=cycle)
        revised_row = [
            f"MP-{cycle.isoformat()}-1-001",
            "MP001",
            message("demo.first_employee"),
            DEPARTMENTS[0],
            first_day,
            message("demo.morning_shift"),
            PUNCH_IN,
            message("demo.explanation.revised"),
            "",
        ]
        send(
            smtp,
            7,
            1,
            period,
            message("demo.email.revised_file", code=f"MP-{cycle.isoformat()}-1-001"),
            [("attendance-department-1-v2.xlsx", workbook([revised_row], {2: [picture("MP001", first_day)]}))],
            cycle_start=cycle,
        )
        send(
            smtp,
            8,
            2,
            period,
            message("demo.email.invalid_template"),
            [
                (
                    "invalid-template.xlsx",
                    workbook(
                        [
                            [
                                "BAD-TEMPLATE",
                                "MP009",
                                message("demo.synthetic_label"),
                                message("demo.invalid_department"),
                                first_day,
                            ]
                        ],
                        bad_template=True,
                    ),
                )
            ],
            cycle_start=cycle,
        )
        send(smtp, 9, 3, period, message("demo.email.missing_workbook"), [], cycle_start=cycle)
        send(
            smtp,
            10,
            None,
            None,
            message("demo.email.unrecognized"),
            [],
            message("demo.email.untrusted_body"),
        )
    # Kept outside classifier inputs; used by the evaluation script only.
    (demo_dir / "answers.json").write_text(json.dumps(answers, ensure_ascii=False, indent=2))
    print("Sent 10 synthetic messages through local SMTP. Case data must be imported by the email pipeline.")
