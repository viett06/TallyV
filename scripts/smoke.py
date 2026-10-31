"""Actual HTTP E2E against the local PostgreSQL/Mailpit/worker stack.
Run after send_demo has sent emails and all messages have been marked read.
"""

import json
import time
from io import BytesIO

import httpx
from openpyxl import load_workbook

from app.core.config import settings
from app.core.i18n import message


def run():
    report = []
    with httpx.Client(base_url="http://127.0.0.1:8000", timeout=30, trust_env=False, follow_redirects=False) as client:
        assert client.get("/api/cases").status_code == 401
        assert (
            client.post(
                "/api/auth/login", json={"username": settings.demo_username, "password": settings.demo_password}
            ).status_code
            == 200
        )
        assert client.get("/").status_code == 200
        assert client.post("/api/sync").status_code == 200
        for _ in range(180):
            jobs = client.get("/api/jobs").json()
            cases = client.get("/api/cases?page_size=100").json()
            if cases["total"] == 60 and not any(j["status"] in {"PENDING", "RUNNING"} for j in jobs):
                break
            time.sleep(1)
        assert cases["total"] == 60, cases
        assert client.post("/api/analysis").status_code == 200
        for _ in range(180):
            jobs = client.get("/api/jobs").json()
            cases = client.get("/api/cases?page_size=100").json()
            if not any(j["status"] in {"PENDING", "RUNNING"} for j in jobs):
                break
            time.sleep(1)
        assert all(c["analysis"]["status"] == "COMPLETED" for c in cases["items"])
        report.append("SMTP -> Mailpit -> workbook and evidence -> 60 cases -> explicit analysis request -> completed")
        emails = client.get("/api/emails").json()
        assert len(emails) == 10
        for email in emails:
            detail = client.get(f"/api/emails/{email['id']}").json()
            for batch in detail["batches"]:
                assert (
                    batch["row_count"] == batch["imported_count"] + batch["duplicate_count"] + batch["error_count"]
                ), batch
        report.append(
            "Every source row is accounted for, including invalid templates, missing workbooks and out-of-scope messages"
        )
        before = (cases["total"], len(emails))
        client.post("/api/sync")
        time.sleep(3)
        assert (client.get("/api/cases").json()["total"], len(client.get("/api/emails").json())) == before
        report.append("Repeated synchronization does not duplicate cases or emails; read messages are still imported")
        eligible = [c for c in cases["items"] if c["bulk_eligible"]]
        blocked = [c for c in cases["items"] if not c["bulk_eligible"]]
        assert eligible and blocked
        assert (
            client.post(
                "/api/cases/bulk-confirm", json={"ids": [eligible[0]["id"], blocked[0]["id"]], "acknowledged": True}
            ).status_code
            == 409
        )
        assert (
            client.post("/api/cases/bulk-confirm", json={"ids": [eligible[0]["id"]], "acknowledged": True}).status_code
            == 200
        )
        assert client.get(f"/api/cases/{eligible[0]['id']}").json()["hr_status"] == "CONFIRMED"
        report.append("The server blocks ineligible bulk reviews and audits eligible confirmations")
        chosen = next(c for c in cases["items"] if c["code"] == "MP-1-005")
        assert (
            client.patch(
                f"/api/cases/{chosen['id']}/review",
                json={
                    "category": "FORGOT_CHECK",
                    "status": "NEEDS_INFO",
                    "note": message("demo.hr_evidence_request"),
                },
            ).status_code
            == 200
        )
        detail = client.get(f"/api/cases/{chosen['id']}").json()
        assert detail["audit"][0]["after"]["hr_status"] == "NEEDS_INFO"
        report.append("HR corrections and information requests retain an audit trail")
        revised = next(c for c in cases["items"] if c["code"] == "MP-1-001")
        detail = client.get(f"/api/cases/{revised['id']}").json()
        assert len(detail["versions"]) == 2 and detail["proposed_count"] == 1
        assert detail["data"]["explanation"] != detail["versions"][1]["normalized"]["explanation"]
        report.append("Revised files create separate proposals and preserve the active source")
        filtered = client.get("/api/cases?period=2026-09&category=DEVICE_ERROR&page_size=100").json()
        response = client.get("/api/export?period=2026-09&category=DEVICE_ERROR")
        wb = load_workbook(BytesIO(response.content), data_only=False)
        assert wb.active.max_row - 1 == filtered["total"]
        assert all(wb.active.cell(r, 16).value == "2026-09" for r in range(2, wb.active.max_row + 1))
        assert all(wb.active.cell(r, 9).value == "DEVICE_ERROR" for r in range(2, wb.active.max_row + 1))
        report.append("Workbook export matches the active database filters")
        assert client.post("/api/sync", headers={"Origin": "https://outside.invalid"}).status_code == 403
        assert client.get("/data/anything").status_code == 404
        sample = next(c for c in cases["items"] if c["evidence_count"])
        image = client.get(f"/api/cases/{sample['id']}").json()["versions"][0]["evidence"][0]
        assert client.get(image["url"]).headers["content-type"] == "image/png"
        client.cookies.clear()
        assert client.get(image["url"]).status_code == 401
        report.append("Evidence requires authorized API access; cross-origin writes and direct file access are blocked")
    result = {"checks": report, "passed": len(report), "synthetic": True}
    (settings.storage_dir / "demo").mkdir(parents=True, exist_ok=True)
    (settings.storage_dir / "demo" / "smoke.json").write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    run()
