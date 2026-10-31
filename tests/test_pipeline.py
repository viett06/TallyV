from datetime import datetime, timezone
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from sqlalchemy import func, select

from app.core.i18n import message
from app.demo.generate import picture, workbook
from app.models import (
    AnalysisResult,
    Attachment,
    AttendanceCase,
    AuditLog,
    CaseVersion,
    Email,
    EvidenceFile,
    ImportBatch,
)
from app.services.analysis import Suggestion, analyze, rule_suggestion
from app.services.email_provider import IncomingEmail
from app.services.excel import parse_workbook
from app.services.imports import process_email, synchronize
from tests.helpers import Provider, ingest, row


def test_partial_import_conserves_rows_and_retry_is_idempotent(context):
    sessions, _ = context
    content = workbook(
        [row(), row("BAD-DATE", day="01/02/2026"), row("BAD-EMP", employee="MP999")],
        {2: [picture("MP001", "2026-10-01")]},
    )
    with sessions() as db:
        email, _ = ingest(db, content)
        batch = db.scalar(select(ImportBatch))
        assert (batch.row_count, batch.imported_count, batch.error_count, batch.duplicate_count) == (3, 1, 2, 0)
        assert email.status == "PARTIAL_ERROR"
        evidence = db.scalar(select(EvidenceFile))
        version = db.get(CaseVersion, evidence.version_id)
        assert version.row_number == 2 and evidence.anchor.startswith("I2")
        assert evidence.ocr_text == ""
        process_email(db, email.id)
        db.commit()
        assert db.scalar(select(func.count()).select_from(AttendanceCase)) == 1
        assert db.scalar(select(func.count()).select_from(EvidenceFile)) == 1
        assert batch.row_count == batch.imported_count + batch.error_count + batch.duplicate_count


def test_repeat_sync_and_duplicate_file_do_not_duplicate_cases(context):
    sessions, _ = context
    content = workbook([row()])
    with sessions() as db:
        ingest(db, content)
        assert synchronize(db, Provider(content)) == {"new": 0, "existing": 1}
        email2, _ = ingest(db, content, "source-2")
        attachment = db.scalar(select(Attachment).where(Attachment.email_id == email2.id))
        assert attachment.status == "DUPLICATE" and attachment.duplicate_of
        assert db.scalar(select(func.count()).select_from(AttendanceCase)) == 1


def test_changed_file_preserves_confirmed_source_and_requires_version_decision(context):
    sessions, client = context
    with sessions() as db:
        ingest(db, workbook([row()]))
        case = db.scalar(select(AttendanceCase))
        old = case.active_version_id
        case.hr_status, case.confirmed_by = "CONFIRMED", "hr.demo"
        db.commit()
        ingest(db, workbook([row(explanation=message("tests.explanation.revised_device"))]), "source-2")
        db.refresh(case)
        versions = list(db.scalars(select(CaseVersion).order_by(CaseVersion.id)))
        assert case.active_version_id == old and len(versions) == 2
        assert case.hr_status == "PENDING_REVIEW" and case.confirmed_by is None
        assert versions[0].raw["explanation"] == message("tests.explanation.forgot_check")
        assert versions[1].status == "PROPOSED"
        assert db.scalar(select(AuditLog)).action == "VERSION_PROPOSED"
        case_id, proposed = case.id, versions[1].id
    response = client.patch(
        f"/api/cases/{case_id}/review",
        json={
            "category": "DEVICE_ERROR",
            "status": "CONFIRMED",
            "acknowledge_exception": True,
            "override_reason": message("tests.exception_reason"),
        },
    )
    assert response.status_code == 409
    assert (
        client.post(
            f"/api/cases/{case_id}/versions/{proposed}",
            json={"action": "APPLY", "reason": message("tests.version_reason")},
        ).status_code
        == 200
    )
    with sessions() as db:
        case = db.get(AttendanceCase, case_id)
        assert case.active_version_id == proposed and case.hr_status == "PENDING_REVIEW"
        assert db.get(CaseVersion, old).status == "SUPERSEDED"


def test_ambiguous_anchor_is_reported_without_guessing(context):
    good = workbook([row()], {2: [picture("MP001", "2026-10-01")]})
    wb = load_workbook(BytesIO(good))
    wb.active._images[0].anchor = TwoCellAnchor(_from=AnchorMarker(col=8, row=1), to=AnchorMarker(col=8, row=2))
    out = BytesIO()
    wb.save(out)
    parsed, errors, count, global_errors = parse_workbook(out.getvalue())
    assert count == 1 and not parsed and errors[0]["row"] == 2 and global_errors


def test_formulas_and_template_errors_are_not_executed(context):
    parsed, errors, count, _ = parse_workbook(workbook([row(explanation='=WEBSERVICE("https://outside.invalid")')]))
    assert not parsed and errors and count == 1
    parsed, errors, count, global_errors = parse_workbook(workbook([row()], bad_template=True))
    assert not parsed and count == len(errors) == 1 and global_errors


def test_bulk_server_enforces_eligibility_and_atomicity(context):
    sessions, client = context
    with sessions() as db:
        ingest(db, workbook([row(), row("TEST-002")]))
        cases = list(db.scalars(select(AttendanceCase).order_by(AttendanceCase.id)))
        analyses = list(db.scalars(select(AnalysisResult).order_by(AnalysisResult.id)))
        analyses[0].status, analyses[0].assessment, analyses[0].category = "COMPLETED", "SUFFICIENT", "FORGOT_CHECK"
        analyses[1].status, analyses[1].assessment, analyses[1].category = "COMPLETED", "MISSING", "FORGOT_CHECK"
        db.commit()
        ids = [c.id for c in cases]
    assert client.post("/api/cases/bulk-confirm", json={"ids": ids, "acknowledged": True}).status_code == 409
    with sessions() as db:
        assert all(c.hr_status == "PENDING_REVIEW" for c in db.scalars(select(AttendanceCase)))
    assert client.post("/api/cases/bulk-confirm", json={"ids": [ids[0]], "acknowledged": True}).json() == {
        "confirmed": 1
    }
    assert client.post("/api/cases/bulk-confirm", json={"ids": [ids[0]], "acknowledged": True}).status_code == 409
    with sessions() as db:
        assert db.scalar(select(AuditLog)).actor == "hr.demo"


def test_hr_correction_audit_and_export_filter_safe_literal(context):
    sessions, client = context
    with sessions() as db:
        ingest(db, workbook([row(), row("TEST-002")]))
        case = db.scalar(select(AttendanceCase).where(AttendanceCase.code == "TEST-001"))
        case_id = case.id
    note = '=HYPERLINK("https://outside.invalid","click")'
    response = client.patch(
        f"/api/cases/{case_id}/review", json={"category": "DEVICE_ERROR", "status": "NEEDS_INFO", "note": note}
    )
    assert response.status_code == 200
    detail = client.get(f"/api/cases/{case_id}").json()
    assert detail["hr_category"] == "DEVICE_ERROR" and detail["audit"][0]["after"]["hr_category"] == "DEVICE_ERROR"
    response = client.get("/api/export?search=TEST-001&hr_status=NEEDS_INFO")
    wb = load_workbook(BytesIO(response.content), data_only=False)
    assert wb.active.max_row == 2
    cell = wb.active.cell(2, 13)
    assert cell.value == note and cell.data_type == "s"
    assert wb.active.cell(2, 11).value == "DEVICE_ERROR"


def test_llm_schema_error_is_visible_and_no_fake_fallback(context, monkeypatch):
    sessions, _ = context
    from app.core.config import settings

    monkeypatch.setattr(settings, "processing_mode", "LOCAL_LLM")
    with sessions() as db:
        ingest(db, workbook([row()]))
        version = db.scalar(select(CaseVersion))
        with patch("app.services.analysis.local_suggestion", side_effect=ValueError("invalid JSON")):
            analyze(db, version.id)
        result = db.scalar(select(AnalysisResult))
        assert result.status == "FAILED" and result.error and result.method == "LOCAL_LLM"
        assert result.category == "UNKNOWN" and result.assessment == "MISSING"


def test_ocr_unavailable_is_explicit_and_does_not_invent_text(context):
    sessions, _ = context
    with sessions() as db:
        ingest(db, workbook([row()], {2: [picture("MP001", "2026-10-01")]}))
        version = db.scalar(select(CaseVersion))
        analyze(db, version.id)
        evidence = db.scalar(select(EvidenceFile))
        result = db.scalar(select(AnalysisResult))
        assert evidence.ocr_status == "UNAVAILABLE" and evidence.ocr_text == ""
        assert result.assessment == "MISSING" and result.category == "FORGOT_CHECK"


def test_auth_files_and_csrf(context):
    sessions, client = context
    with sessions() as db:
        ingest(db, workbook([row()], {2: [picture("MP001", "2026-10-01")]}))
    assert client.post("/api/sync", headers={"Origin": "https://outside.invalid"}).status_code == 403
    assert client.get("/data/anything").status_code == 404
    client.cookies.clear()
    for path in ["/api/cases", "/api/emails", "/api/export", "/api/evidence/1", "/api/attachments/1"]:
        assert client.get(path).status_code == 401


def test_rule_classifier_schema_and_conflicting_causes():
    assert (
        Suggestion.model_validate(rule_suggestion(message("tests.explanation.forgot_check"))).category == "FORGOT_CHECK"
    )
    assert rule_suggestion(message("tests.explanation.multiple_causes")).category == "UNKNOWN"
    assert rule_suggestion(message("tests.explanation.overtime")).category == "UNKNOWN"


def test_no_external_endpoint_and_no_html_execution(context):
    import pytest
    from pydantic import ValidationError

    from app.core.config import Settings

    with pytest.raises(ValidationError):
        Settings(mailpit_url="https://public.example.com")
    with pytest.raises(ValidationError):
        Settings(ollama_url="http://public.example.com")
    sessions, client = context

    class UnsafeProvider:
        def messages(self):
            return [
                IncomingEmail(
                    "html-email",
                    "<html-email@demo.invalid>",
                    "unknown@demo.invalid",
                    "hr@demo.invalid",
                    message("tests.email_outside_scope"),
                    "<script>alert(1)</script> https://outside.invalid",
                    datetime.now(timezone.utc),
                    [],
                )
            ]

    with sessions() as db:
        synchronize(db, UnsafeProvider())
        db.commit()
        email = db.scalar(select(Email))
        email_id = email.id
        assert email.status == "NEEDS_REVIEW"
    response = client.get(f"/api/emails/{email_id}")
    assert response.headers["content-type"] == "application/json"
    assert response.json()["body"].startswith("<script>")  # JSON data; UI uses React text nodes.
    assert "connect-src 'self'" in response.headers["content-security-policy"]
    assert client.post(f"/api/emails/{email_id}/retry").status_code == 422


def test_local_llm_payload_excludes_employee_name_and_validates_json(context, monkeypatch):
    from app.services.analysis import local_suggestion

    class FakeResponse:
        def raise_for_status(self):
            pass

        def json(self):
            return {
                "message": {
                    "content": '{"category":"FORGOT_CHECK","reason":"The explanation reports a forgotten attendance check","extracted":{},"missing":[],"conflicts":[]}'
                }
            }

    class FakeClient:
        def __init__(self, **kwargs):
            assert kwargs["trust_env"] is False and kwargs["follow_redirects"] is False

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def post(self, url, json):
            assert set(__import__("json").loads(json["messages"][1]["content"])) == {
                "case_code",
                "explanation",
                "evidence",
            }
            assert json["format"]["additionalProperties"] is False
            return FakeResponse()

    monkeypatch.setattr("app.services.analysis.httpx.Client", FakeClient)
    result = local_suggestion("TEST-001", message("category.forgot_check"), [])
    assert result.category == "FORGOT_CHECK"


def test_ocr_normalization_preserves_raw_text_and_detects_conflict(context, monkeypatch):
    import pytesseract

    from app.core.config import settings

    sessions, _ = context
    monkeypatch.setattr(settings, "ocr_enabled", True)
    raw = "EMPLOYEE: MPOO1\nDATE: 2026-10-02\nSTATUS: RECORDED\n"
    monkeypatch.setattr(pytesseract, "image_to_string", lambda *a, **k: raw)
    with sessions() as db:
        ingest(db, workbook([row()], {2: [picture("MP001", "2026-10-01")]}))
        version = db.scalar(select(CaseVersion))
        analyze(db, version.id)
        evidence = db.scalar(select(EvidenceFile))
        result = db.scalar(select(AnalysisResult))
        assert evidence.ocr_text == raw and evidence.extracted["employee_id"] == "MP001"
        assert "normalization" in evidence.extracted
        assert result.assessment == "CONFLICT" and any(message("excel.header.date") in c for c in result.conflicts)


def test_worker_retry_rolls_back_partial_job_and_does_not_duplicate(context, monkeypatch):
    from datetime import timedelta

    from app import worker
    from app.core.time import utcnow
    from app.models import ProcessingJob

    sessions, _ = context
    monkeypatch.setattr(worker, "SessionLocal", sessions)
    content = workbook([row()], {2: [picture("MP001", "2026-10-01")]})
    with sessions() as db:
        synchronize(db, Provider(content))
        db.commit()
    actual = worker.process_email

    def fail_after_processing(db, email_id):
        actual(db, email_id)
        raise RuntimeError("simulated crash before transaction commit")

    monkeypatch.setattr(worker, "process_email", fail_after_processing)
    assert worker.run_one()
    with sessions() as db:
        assert db.scalar(select(func.count()).select_from(AttendanceCase)) == 0
        job = db.scalar(select(ProcessingJob))
        assert job.status == "PENDING" and job.attempts == 1
        job.available_at = utcnow() - timedelta(seconds=1)
        db.commit()
    monkeypatch.setattr(worker, "process_email", actual)
    assert worker.run_one()
    with sessions() as db:
        assert db.scalar(select(func.count()).select_from(AttendanceCase)) == 1
        assert db.scalar(select(func.count()).select_from(EvidenceFile)) == 1
        job = db.scalar(select(ProcessingJob).where(ProcessingJob.kind == "IMPORT_EMAIL"))
        assert job.status == "DONE" and job.attempts == 2


def test_rule_change_invalidates_previous_confirmation(context):
    sessions, client = context
    with sessions() as db:
        ingest(db, workbook([row()]))
        case = db.scalar(select(AttendanceCase))
        result = db.scalar(select(AnalysisResult))
        result.category = "FORGOT_CHECK"
        case.hr_status = "CONFIRMED"
        case.confirmed_by = "hr.demo"
        case_id = case.id
        db.commit()
    response = client.put(
        "/api/rules/FORGOT_CHECK",
        json={"require_evidence": True, "require_readable_ocr": True, "require_shift": True, "require_punch": True},
    )
    assert response.status_code == 200
    with sessions() as db:
        case = db.get(AttendanceCase, case_id)
        result = db.scalar(select(AnalysisResult))
        assert case.hr_status == "PENDING_REVIEW" and case.confirmed_by is None
        assert result.status == "NOT_STARTED"
        assert db.scalar(select(AuditLog).where(AuditLog.case_id == case_id)).action == "RULE_CHANGED"


def test_watermark_only_ocr_is_not_readable_business_evidence(context, monkeypatch):
    import pytesseract

    from app.core.config import settings

    sessions, _ = context
    monkeypatch.setattr(settings, "ocr_enabled", True)
    monkeypatch.setattr(pytesseract, "image_to_string", lambda *a, **k: message("tests.ocr.watermark_only"))
    with sessions() as db:
        ingest(db, workbook([row()], {2: [picture("MP001", "2026-10-01", unreadable=True)]}))
        analyze(db, db.scalar(select(CaseVersion.id)))
        evidence = db.scalar(select(EvidenceFile))
        assert evidence.ocr_status == "UNREADABLE"
        assert evidence.ocr_text.startswith(message("tests.ocr.watermark_prefix"))


def test_fully_invalid_template_is_failed_not_completed_or_partial(context):
    sessions, _ = context
    with sessions() as db:
        email, _ = ingest(db, workbook([row()], bad_template=True))
        assert email.status == "FAILED"
        assert db.scalar(select(Attachment)).status == "FAILED"
        assert db.scalar(select(func.count()).select_from(AttendanceCase)) == 0
        batch = db.scalar(select(ImportBatch))
        assert batch.row_count == batch.error_count == 1
